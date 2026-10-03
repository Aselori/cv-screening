"""Pruebas sobre el conjunto propio en español (data/samples/es)."""

import csv
import json
import re
from pathlib import Path

import docx
import pdfplumber
import pytest

from cv_screening.anonymization import anonymize
from cv_screening.schemas import FitLabel, Vacancy

SAMPLES = Path(__file__).resolve().parents[1] / "data" / "samples" / "es"
SOURCES = sorted((SAMPLES / "cvs").glob("cv-*.md"))


def _docx_text(path: Path) -> str:
    document = docx.Document(path)
    cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
    return "\n".join(cells + [p.text for p in document.paragraphs])


def _pdf_text(path: Path) -> str:
    with pdfplumber.open(path) as pdf:
        return "\n".join(page.extract_text() or "" for page in pdf.pages)


def _personal_data(source: str) -> dict[str, str]:
    match = re.search(r"^# (.+)$", source, re.M) or re.search(r"Nombre: (.+)$", source, re.M)
    name = match.group(1) if match else source.splitlines()[1]
    return {
        "nombre": name.title(),
        "nombre de pila": name.title().split()[0],
        "correo": re.search(r"[\w.]+@example\.com", source).group(0),
        "teléfono": re.search(r"5555[ -]?0\d{3}|8155550\d{3}", source).group(0)[-4:],
    }


def test_there_are_24_sources_with_docx_and_pdf():
    assert len(SOURCES) == 24
    for source in SOURCES:
        assert (SAMPLES / "docx" / f"{source.stem}.docx").exists()
        assert (SAMPLES / "pdf" / f"{source.stem}.pdf").exists()


@pytest.mark.parametrize("source", SOURCES, ids=lambda p: p.stem)
@pytest.mark.parametrize("fmt", ["docx", "pdf"])
def test_anonymization_removes_personal_data(source: Path, fmt: str):
    extract = _docx_text if fmt == "docx" else _pdf_text
    text = anonymize(extract(SAMPLES / fmt / f"{source.stem}.{fmt}")).text
    for kind, value in _personal_data(source.read_text(encoding="utf-8")).items():
        assert not re.search(rf"(?<!\w){re.escape(value)}(?!\w)", text, re.I), kind


def test_vacancies_follow_the_schema():
    vacancies = json.loads((SAMPLES / "vacancies.json").read_text(encoding="utf-8"))
    assert [Vacancy.model_validate(v).vacancy_id for v in vacancies] == [
        "v-es-01",
        "v-es-02",
        "v-es-03",
        "v-es-04",
    ]


def test_labels_cover_every_cv_vacancy_pair():
    with open(SAMPLES / "labels.csv", encoding="utf-8") as f:
        rows = list(csv.DictReader(f))
    pairs = {(r["cv_id"], r["vacancy_id"]) for r in rows}
    assert len(rows) == len(pairs) == 24 * 4
    assert {r["cv_id"] for r in rows} == {s.stem for s in SOURCES}
    assert all(FitLabel(r["label"]) for r in rows)

"""Genera los CVs de muestra en DOCX y PDF a partir de sus fuentes en Markdown.

Uso (desde la raíz del repositorio):
    python scripts/build_samples.py

Requiere LibreOffice (`soffice`) para convertir DOCX a PDF. Los archivos generados se
guardan en Git, así que solo hace falta ejecutarlo al cambiar las fuentes.

Cada CV usa uno de tres formatos, para que el parser se pruebe con documentos variados:
    0: estilos de título de Word
    1: texto normal con títulos en negritas y mayúsculas
    2: encabezado en una tabla de dos columnas (nombre a la izquierda, contacto a la derecha)
"""

import re
import shutil
import subprocess
import sys
from dataclasses import dataclass, field
from pathlib import Path

from docx import Document
from docx.shared import Pt

SAMPLES_DIR = Path(__file__).resolve().parents[1] / "data" / "samples" / "es"
SOURCE_DIR = SAMPLES_DIR / "cvs"
DOCX_DIR = SAMPLES_DIR / "docx"
PDF_DIR = SAMPLES_DIR / "pdf"


@dataclass
class Section:
    title: str
    items: list[tuple[str, str]] = field(default_factory=list)  # (tipo, texto)


@dataclass
class SourceCV:
    name: str | None
    header: list[str]
    sections: list[Section]


def parse_source(text: str) -> SourceCV:
    """Lee el Markdown simplificado de las fuentes (#, ##, ###, viñetas y párrafos)."""
    name, header, sections = None, [], []
    for line in text.splitlines():
        line = line.rstrip()
        if not line:
            continue
        if line.startswith("## "):
            sections.append(Section(line[3:]))
        elif line.startswith("# ") and not sections:
            name = line[2:]
        elif not sections:
            header.append(line)
        elif line.startswith("### "):
            sections[-1].items.append(("entry", line[4:]))
        elif line.startswith("- "):
            sections[-1].items.append(("bullet", line[2:]))
        else:
            sections[-1].items.append(("paragraph", line))
    return SourceCV(name, header, sections)


def _bold(document: Document, text: str, size: int | None = None):
    run = document.add_paragraph().add_run(text)
    run.bold = True
    if size:
        run.font.size = Pt(size)


def _add_sections(document: Document, cv: SourceCV, plain: bool) -> None:
    for section in cv.sections:
        if plain:
            _bold(document, section.title.upper(), 12)
        else:
            document.add_heading(section.title, level=2)
        for kind, text in section.items:
            if kind == "entry":
                if plain:
                    _bold(document, text)
                else:
                    document.add_heading(text, level=3)
            elif kind == "bullet":
                if plain:
                    document.add_paragraph(f"• {text}")
                else:
                    document.add_paragraph(text, style="List Bullet")
            else:
                document.add_paragraph(text)


def build_docx(cv: SourceCV, layout: int, path: Path) -> None:
    document = Document()
    if layout == 2 and cv.name:
        table = document.add_table(rows=1, cols=2)
        left, right = table.rows[0].cells
        left.paragraphs[0].add_run(cv.name).bold = True
        for line in cv.header[:1]:
            left.add_paragraph(line)
        right.paragraphs[0].text = cv.header[1] if len(cv.header) > 1 else ""
        for line in cv.header[2:]:
            right.add_paragraph(line)
    else:
        if cv.name:
            if layout == 0:
                document.add_heading(cv.name, level=1)
            else:
                _bold(document, cv.name, 16)
        for line in cv.header:
            document.add_paragraph(line)
    _add_sections(document, cv, plain=layout == 1)
    document.save(path)


def main() -> None:
    soffice = shutil.which("soffice")
    if soffice is None:
        sys.exit("No se encontró LibreOffice (soffice); es necesario para generar los PDF.")
    DOCX_DIR.mkdir(exist_ok=True)
    PDF_DIR.mkdir(exist_ok=True)
    sources = sorted(SOURCE_DIR.glob("cv-*.md"))
    for source in sources:
        number = int(re.search(r"\d+", source.stem).group())
        cv = parse_source(source.read_text(encoding="utf-8"))
        build_docx(cv, layout=number % 3, path=DOCX_DIR / f"{source.stem}.docx")
    subprocess.run(
        [soffice, "--headless", "--convert-to", "pdf", "--outdir", str(PDF_DIR)]
        + [str(DOCX_DIR / f"{s.stem}.docx") for s in sources],
        check=True,
        capture_output=True,
    )
    print(f"{len(sources)} CVs generados en {DOCX_DIR} y {PDF_DIR}")


if __name__ == "__main__":
    main()
